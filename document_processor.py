"""
Document processing module
Handles PDF, DOCX, TXT, and XLSX files
"""

import os
from pathlib import Path
from typing import List, Dict
import pypdf
from docx import Document
import openpyxl
import pandas as pd

class DocumentProcessor:
    """Extract text from various document formats"""
    
    SUPPORTED_FORMATS = {'.pdf', '.docx', '.txt', '.xlsx'}
    
    def __init__(self, folder_path: str = "watched_folder"):
        self.folder_path = Path(folder_path)
        self.folder_path.mkdir(exist_ok=True)
    
    def get_all_documents(self) -> List[Path]:
        """Get all supported documents in watched folder"""
        docs = []
        for file in self.folder_path.iterdir():
            if file.suffix.lower() in self.SUPPORTED_FORMATS:
                docs.append(file)
        return docs
    
    def extract_text(self, file_path: Path) -> Dict[str, str]:
        """
        Extract text from a document
        Returns: {'filename': str, 'content': str, 'source': str}
        """
        suffix = file_path.suffix.lower()
        
        try:
            if suffix == '.pdf':
                text = self._extract_pdf(file_path)
            elif suffix == '.docx':
                text = self._extract_docx(file_path)
            elif suffix == '.txt':
                text = self._extract_txt(file_path)
            elif suffix == '.xlsx':
                text = self._extract_xlsx(file_path)
            else:
                return None
            
            return {
                'filename': file_path.name,
                'content': text,
                'source': str(file_path)
            }
        except Exception as e:
            print(f"Error processing {file_path.name}: {e}")
            return None
    
    def _extract_pdf(self, file_path: Path) -> str:
        """Extract text from PDF"""
        text = []
        with open(file_path, 'rb') as f:
            pdf = pypdf.PdfReader(f)
            for page in pdf.pages:
                text.append(page.extract_text())
        return '\n'.join(text)
    
    def _extract_docx(self, file_path: Path) -> str:
        """Extract text from DOCX"""
        doc = Document(file_path)
        text = []
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text)
        return '\n'.join(text)
    
    def _extract_txt(self, file_path: Path) -> str:
        """Extract text from TXT"""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    
    def _extract_xlsx(self, file_path: Path) -> str:
        """Extract text from Excel (all sheets)"""
        text = []
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text.append(f"\n--- Sheet: {sheet_name} ---\n")
            
            # Convert to pandas for easier text extraction
            data = []
            for row in sheet.iter_rows(values_only=True):
                data.append(row)
            
            if data:
                df = pd.DataFrame(data[1:], columns=data[0])
                text.append(df.to_string(index=False))
        
        return '\n'.join(text)

if __name__ == "__main__":
    # Test the processor
    processor = DocumentProcessor()
    docs = processor.get_all_documents()
    print(f"Found {len(docs)} documents")
    
    for doc in docs:
        result = processor.extract_text(doc)
        if result:
            print(f"\n{result['filename']}: {len(result['content'])} chars")